#!/usr/bin/env python3
"""Cross-check the generated specs against the source workbook.

`parse_source.py` and `source_spec.py` can only be trusted if every label in the
workbook ends up somewhere: captured as a field, deliberately taken from the
OpenMRS patient record, or explicitly skipped. This script fails loudly if a
workbook label is unaccounted for, which is the failure mode that let the whole
Back tab, the Delivery Register and Family Planning go missing the first time.

    python3 verify_source_coverage.py
"""
import os
import sys

import openpyxl
from openpyxl.utils import get_column_letter

import source_spec as S

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.join(HERE, "source", "mother-and-neonate-register.xlsx")
BACK = "Mother & Neonate (Back)"
FAMILY = "Family Planning"
DELIVERY = "Delivery Register"
FRONT = "Mother & Neonate (Front)"

failures = []
notes = []


def check(condition, message):
    if not condition:
        failures.append(message)


def norm(text):
    return " ".join(str(text or "").split()).strip().lower()


def main():
    wb = openpyxl.load_workbook(BOOK, data_only=True)

    mother = S.mother_fields()
    pn = S.postnatal_fields()
    fp = S.family_planning_fields()
    dr = S.delivery_fields()

    # ---- Back tab: every printed code must exist in the spec --------------
    ws = wb[BACK]
    src_codes = set()
    for row in range(1, ws.max_row + 1):
        for col in (1, 17):
            v = norm(ws.cell(row, col).value)
            if v and (v[0] == "d" and v[1:].isdigit()) or (v.startswith(("pr", "nn"))
                                                           and v[2:].isdigit()):
                src_codes.add(v.upper())

    spec_codes = {f["source_code"].upper() for f in mother if "source_code" in f}
    spec_codes |= {f["code"].upper() for f in pn["postnatal_mother"]["fields"]}
    spec_codes |= {f["code"].upper() for f in pn["postnatal_baby"]["fields"]}
    # prompt sub-fields are synthesised (M1v1) and have no source code
    spec_codes = {c for c in spec_codes if not c.endswith("V1") and "V" not in c}

    missing = sorted(src_codes - spec_codes)
    check(not missing, f"Back tab codes missing from spec: {missing}")

    # ---- Back tab: visit windows -----------------------------------------
    parsed_windows = S._load()[BACK]["visit_windows"]
    check(parsed_windows == S.VISIT_WINDOWS,
          f"visit windows drifted: workbook {parsed_windows} vs spec {S.VISIT_WINDOWS}")

    # ---- Family Planning: every printed column accounted for -------------
    # Compared against the spec's own record of what it consumed, so the two
    # cannot drift; corrections are reported rather than silently applied.
    ws = wb[FAMILY]
    consumed = S.family_planning_consumed()
    known = {norm(S.fix(x)) for x in (consumed["labels"] + consumed["options"]
                                        + consumed["subgroups"]
                                        + consumed["from_patient"]
                                        + consumed["skipped"])}
    for col in range(1, ws.max_column + 1):
        for row in (3, 4, 5):
            raw = ws.cell(row, col).value
            printed = " ".join(str(raw or "").split()).strip()
            if not printed or printed == "(please tick)":
                continue
            # A printed label may be typo-corrected and may carry a trailing
            # "( Y / N )" marker, so any normalised variant of it is accepted.
            variants = {v for v in S._label_keys(printed) if v}
            if variants & known or norm(S.fix(printed)) in known:
                if norm(S.fix(printed)) != norm(printed):
                    notes.append(f"Family Planning source typo corrected: "
                                 f"{printed!r} -> {S.fix(printed)!r}")
                continue
            failures.append(
                f"Family Planning cell {get_column_letter(col)}{row} "
                f"unaccounted for: {printed!r}")

    # ---- Delivery Register: every non-empty header cell accounted for -----
    ws = wb[DELIVERY]
    consumed = {(r, c) for (r, c) in S.DR_CONSUMED}
    # Row-3 group headings are carried as each field's "group" rather than as a
    # label, so treat a header cell as consumed when it names a declared group.
    declared_groups = {norm(g) for g in S.GROUP_HEADERS}
    for row in (3,):
        for col in range(1, ws.max_column + 1):
            if norm(ws.cell(row, col).value) in declared_groups:
                consumed.add((row, col))
    patient = set(S.DR_FROM_PATIENT)
    skipped = set(S.DR_SKIP)
    for row in (3, 4, 5):
        for col in range(1, ws.max_column + 1):
            v = norm(ws.cell(row, col).value)
            if not v:
                continue
            if (row, col) in consumed or (row, col) in patient or (row, col) in skipped:
                continue
            failures.append(
                f"Delivery Register cell {get_column_letter(col)}{row} unaccounted for: {v!r}")
    for g in sorted(S.GROUP_HEADERS):
        if norm(g) not in {norm(f["group"]) for f in dr if f.get("group")}:
            failures.append(f"Delivery Register declared group unused: {g!r}")

    # ---- Front tab: row 4 must be fully classified -----------------------
    # The row-4 labels are all printed as "Client Name:" etc., so the trailing
    # colon is dropped before comparing against the classification lists.
    ws = wb[FRONT]
    row4 = [str(ws.cell(4, c).value or "").strip().rstrip(":").strip()
            for c in range(1, ws.max_column + 1)]
    row4 = [v for v in row4 if v]
    for label in row4:
        if label in S.FRONT_FROM_PATIENT:
            continue
        if label in S.FRONT_CAPTURED:
            continue
        failures.append(f"Front tab row 4 label unclassified: {label!r}")
    for label in S.FRONT_CAPTURED:
        if label not in row4:
            failures.append(f"Front tab captured label absent from workbook: {label!r}")
    notes.append(f"Front tab row 4: from patient {S.FRONT_FROM_PATIENT}, "
                 f"captured {S.FRONT_CAPTURED}")

    # ---- Row 1 titles and row 2 facility banners on every tab -------------
    # The banner components differ between tabs (the Front and Back MNR sheets
    # print no Ownership line, the Delivery Register prints no In Facility /
    # Outreach), so each tab is checked against the components it actually has
    # and the difference is reported rather than assumed uniform.
    seen = {c: set() for c in S.FACILITY_HEADER_COMPONENTS}
    for sheet in (FRONT, BACK, FAMILY, DELIVERY):
        ws = wb[sheet]
        check(bool(norm(ws.cell(1, 1).value)), f"{sheet}: row 1 has no register title")
        banner = norm(ws.cell(2, 1).value)
        if not banner:
            failures.append(f"{sheet}: row 2 facility banner is empty")
            continue
        present = {c for c in S.FACILITY_HEADER_COMPONENTS if c in banner}
        for c in present:
            seen[c].add(sheet)
        missing = [c for c in S.FACILITY_HEADER_COMPONENTS if c not in present]
        notes.append(f"{sheet}: row 2 banner -> encounter/location metadata; "
                     f"not printed on this tab: {missing}")
    for component, sheets in sorted(seen.items()):
        if not sheets:
            failures.append(
                f"Facility header component {component!r} mapped to encounter "
                f"metadata but never appears in any tab")

    # ---- patient-sourced columns must not also be captured as form fields ---
    # Coverage alone would not notice this: demoting a patient-sourced column to a
    # captured field still "accounts for" it, but duplicates data the OpenMRS
    # patient header already holds and lets the two disagree. The required set is
    # asserted too, so quietly dropping a column from the policy is caught.
    for name, fields, policy, required in (
            ("Family Planning", fp, S.FP_FROM_PATIENT, S.FP_PATIENT_REQUIRED),
            ("Delivery Register", dr,
             [S.delivery_label(c) for c in S.DR_FROM_PATIENT],
             S.DR_PATIENT_REQUIRED)):
        present = {norm(x) for x in policy}
        for label in required:
            if norm(label) not in present:
                failures.append(
                    f"{name}: {label!r} must be read from the patient record "
                    f"but is no longer in the patient-sourced list")
        for f in fields:
            if norm(f["label"]) in present:
                failures.append(
                    f"{name}: {f['code']} {f['label']!r} is read from the patient "
                    f"record but also captured as a form field")

    # ---- report ----------------------------------------------------------
    print(f"Back tab codes in workbook : {len(src_codes)}")
    print(f"M1-M20 fields              : {len(mother)}")
    print(f"PR fields (x3)             : {len(pn['postnatal_mother']['fields'])}")
    print(f"NN fields (x3)             : {len(pn['postnatal_baby']['fields'])}")
    print(f"FP fields                  : {len(fp)}")
    print(f"DR fields                  : {len(dr)}")
    for n in notes:
        print(f"note: {n}")
    if failures:
        print(f"\nFAIL: {len(failures)} unaccounted workbook label(s)")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nOK: every workbook label is accounted for")
    return 0


if __name__ == "__main__":
    sys.exit(main())
